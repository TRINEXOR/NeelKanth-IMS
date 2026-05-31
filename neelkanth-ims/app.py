from flask import Flask, render_template, request, jsonify, send_file, Response, stream_with_context
import sqlite3, os, csv, io, json, time, sys, platform
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
DATA_DIR = os.environ.get('DATA_DIR', os.path.dirname(__file__))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, 'inventory.db')

# ─────────────────────────────────────────────
# Database
# ─────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS items (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
            category TEXT NOT NULL DEFAULT 'General', quantity INTEGER NOT NULL DEFAULT 0,
            unit TEXT NOT NULL DEFAULT 'pcs', low_stock INTEGER NOT NULL DEFAULT 10,
            price REAL NOT NULL DEFAULT 0.0, supplier TEXT, notes TEXT, updated_at TEXT NOT NULL)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS activity_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER,
            item_name TEXT, action TEXT, details TEXT, timestamp TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS alert_config (
            id INTEGER PRIMARY KEY CHECK (id=1),
            enabled INTEGER NOT NULL DEFAULT 1,
            last_notified TEXT)''')
        conn.execute("INSERT OR IGNORE INTO alert_config (id,enabled) VALUES (1,1)")
        _migrate_alert_config(conn)
        cur = conn.execute('SELECT COUNT(*) FROM items')
        if cur.fetchone()[0] == 0:
            seed = [
                ('Rice (Basmati)','Grains',    120,'kg',  20,2.50,'FoodCorp',  'Long grain'),
                ('Wheat Flour',  'Grains',       8,'kg',  15,1.80,'MillersInc','All-purpose'),
                ('Cooking Oil',  'Oils',          5,'L',   10,3.20,'PureOils',  None),
                ('Canned Tomato','Canned Goods', 42,'cans',15,0.90,'CanCo',     None),
                ('Lentils',      'Legumes',       3,'kg',  10,2.10,'FoodCorp',  'Red lentils'),
                ('Sugar',        'Condiments',   30,'kg',  10,1.20,'SugarMill', None),
                ('Salt',         'Condiments',   25,'kg',   5,0.40,'SaltWorks', 'Iodised'),
                ('Pasta',        'Grains',        7,'packs',10,1.60,'PastaCo',  '500g packs'),
                ('Baby Formula', 'Baby',          2,'cans', 5,12.00,'NutriKids','Stage 2'),
                ('Soap Bars',    'Hygiene',      18,'bars', 8,0.70,'CleanCo',   None),
            ]
            now = datetime.now().isoformat(timespec='seconds')
            for r in seed:
                conn.execute('INSERT INTO items (name,category,quantity,unit,low_stock,price,supplier,notes,updated_at) VALUES (?,?,?,?,?,?,?,?,?)', (*r, now))
        conn.commit()

def _migrate_alert_config(conn):
    """Drop legacy email columns; keep enabled + last_notified only."""
    cols = {r[1] for r in conn.execute('PRAGMA table_info(alert_config)').fetchall()}
    if cols == {'id', 'enabled', 'last_notified'}:
        return
    if cols == {'id', 'enabled'}:
        conn.execute('ALTER TABLE alert_config ADD COLUMN last_notified TEXT')
        return
    row = conn.execute('SELECT * FROM alert_config WHERE id=1').fetchone()
    enabled = row['enabled'] if row else 1
    last_notified = None
    if row:
        keys = row.keys()
        last_notified = row['last_notified'] if 'last_notified' in keys else (
            row['last_sent'] if 'last_sent' in keys else None)
    conn.execute('''CREATE TABLE alert_config_new (
        id INTEGER PRIMARY KEY CHECK (id=1),
        enabled INTEGER NOT NULL DEFAULT 1,
        last_notified TEXT)''')
    conn.execute('INSERT INTO alert_config_new (id,enabled,last_notified) VALUES (?,?,?)',
                 (1, enabled, last_notified))
    conn.execute('DROP TABLE alert_config')
    conn.execute('ALTER TABLE alert_config_new RENAME TO alert_config')

# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────
def get_low_stock_items():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM items WHERE quantity <= low_stock ORDER BY quantity ASC').fetchall()
    return [dict(r) for r in rows]

# ─────────────────────────────────────────────
# Core CRUD (unchanged)
# ─────────────────────────────────────────────
@app.route('/')
def index(): return render_template('index.html')

@app.route('/api/items')
def get_items():
    search = request.args.get('search','').strip()
    category = request.args.get('category','').strip()
    low_only = request.args.get('low_only','false') == 'true'
    sort_by = request.args.get('sort','name')
    order = request.args.get('order','asc')
    allowed = {'name','category','quantity','price','updated_at'}
    if sort_by not in allowed: sort_by = 'name'
    order = 'DESC' if order == 'desc' else 'ASC'
    q = 'SELECT * FROM items WHERE 1=1'; p = []
    if search:
        q += ' AND (name LIKE ? OR supplier LIKE ? OR notes LIKE ?)'
        p += [f'%{search}%']*3
    if category: q += ' AND category = ?'; p.append(category)
    if low_only: q += ' AND quantity <= low_stock'
    q += f' ORDER BY {sort_by} {order}'
    with get_db() as conn:
        rows = conn.execute(q, p).fetchall()
    items = []
    for r in rows:
        d = dict(r); d['is_low'] = d['quantity'] <= d['low_stock']; d['is_out'] = d['quantity'] == 0
        items.append(d)
    return jsonify(items)

@app.route('/api/stats')
def get_stats():
    with get_db() as conn:
        total     = conn.execute('SELECT COUNT(*) FROM items').fetchone()[0]
        low_stock = conn.execute('SELECT COUNT(*) FROM items WHERE quantity<=low_stock AND quantity>0').fetchone()[0]
        out_stock = conn.execute('SELECT COUNT(*) FROM items WHERE quantity=0').fetchone()[0]
        value     = conn.execute('SELECT COALESCE(SUM(quantity*price),0) FROM items').fetchone()[0]
        cats      = conn.execute('SELECT COUNT(DISTINCT category) FROM items').fetchone()[0]
    return jsonify({'total':total,'low_stock':low_stock,'out_stock':out_stock,'total_value':round(value,2),'categories':cats})

@app.route('/api/categories')
def get_categories():
    with get_db() as conn:
        rows = conn.execute('SELECT DISTINCT category FROM items ORDER BY category').fetchall()
    return jsonify([r[0] for r in rows])

@app.route('/api/log')
def get_log():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM activity_log ORDER BY id DESC LIMIT 50').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/items', methods=['POST'])
def create_item():
    data = request.get_json()
    for f in ['name','category','quantity','unit','low_stock','price']:
        if f not in data: return jsonify({'error': f'Missing: {f}'}), 400
    now = datetime.now().isoformat(timespec='seconds')
    with get_db() as conn:
        cur = conn.execute('INSERT INTO items (name,category,quantity,unit,low_stock,price,supplier,notes,updated_at) VALUES (?,?,?,?,?,?,?,?,?)',
            (data['name'],data['category'],int(data['quantity']),data['unit'],int(data['low_stock']),float(data['price']),data.get('supplier',''),data.get('notes',''),now))
        conn.execute('INSERT INTO activity_log VALUES (NULL,?,?,?,?,?)', (cur.lastrowid,data['name'],'Added',f"Qty: {data['quantity']} {data['unit']}",now))
        conn.commit()
    return jsonify({'id':cur.lastrowid,'message':'Item created'}), 201

@app.route('/api/items/<int:item_id>', methods=['PUT'])
def update_item(item_id):
    data = request.get_json(); now = datetime.now().isoformat(timespec='seconds')
    with get_db() as conn:
        e = conn.execute('SELECT * FROM items WHERE id=?',(item_id,)).fetchone()
        if not e: return jsonify({'error':'Not found'}), 404
        conn.execute('UPDATE items SET name=?,category=?,quantity=?,unit=?,low_stock=?,price=?,supplier=?,notes=?,updated_at=? WHERE id=?',
            (data.get('name',e['name']),data.get('category',e['category']),int(data.get('quantity',e['quantity'])),
             data.get('unit',e['unit']),int(data.get('low_stock',e['low_stock'])),float(data.get('price',e['price'])),
             data.get('supplier',e['supplier']),data.get('notes',e['notes']),now,item_id))
        conn.execute('INSERT INTO activity_log VALUES (NULL,?,?,?,?,?)',
            (item_id,data.get('name',e['name']),'Updated',f"Qty: {data.get('quantity',e['quantity'])}",now))
        conn.commit()
    return jsonify({'message':'Item updated'})

@app.route('/api/items/<int:item_id>', methods=['DELETE'])
def delete_item(item_id):
    now = datetime.now().isoformat(timespec='seconds')
    with get_db() as conn:
        e = conn.execute('SELECT * FROM items WHERE id=?',(item_id,)).fetchone()
        if not e: return jsonify({'error':'Not found'}), 404
        conn.execute('DELETE FROM items WHERE id=?',(item_id,))
        conn.execute('INSERT INTO activity_log VALUES (NULL,?,?,?,?,?)',(item_id,e['name'],'Deleted','—',now))
        conn.commit()
    return jsonify({'message':'Item deleted'})

# ─────────────────────────────────────────────
# NEW: Analytics API
# ─────────────────────────────────────────────
@app.route('/api/analytics')
def get_analytics():
    with get_db() as conn:
        by_cat = conn.execute('''
            SELECT category,
                   COUNT(*) as item_count,
                   SUM(quantity) as total_qty,
                   ROUND(SUM(quantity*price),2) as stock_value
            FROM items GROUP BY category ORDER BY stock_value DESC''').fetchall()
        top_value = conn.execute('''
            SELECT name, category, quantity, price, ROUND(quantity*price,2) as total_value
            FROM items ORDER BY total_value DESC LIMIT 5''').fetchall()
        health = conn.execute('''
            SELECT
                SUM(CASE WHEN quantity=0 THEN 1 ELSE 0 END) as out_of_stock,
                SUM(CASE WHEN quantity>0 AND quantity<=low_stock THEN 1 ELSE 0 END) as low,
                SUM(CASE WHEN quantity>low_stock AND quantity<=low_stock*2 THEN 1 ELSE 0 END) as moderate,
                SUM(CASE WHEN quantity>low_stock*2 THEN 1 ELSE 0 END) as healthy
            FROM items''').fetchone()
        restock = conn.execute('''
            SELECT name, category, quantity, low_stock, unit,
                   ROUND((CAST(quantity AS REAL)/low_stock)*100,1) as pct
            FROM items WHERE quantity <= low_stock ORDER BY pct ASC LIMIT 8''').fetchall()
        trend = conn.execute('''
            SELECT DATE(timestamp) as day, COUNT(*) as ops
            FROM activity_log
            WHERE timestamp >= DATE('now','-14 days')
            GROUP BY day ORDER BY day''').fetchall()
    return jsonify({
        'by_category': [dict(r) for r in by_cat],
        'top_value_items': [dict(r) for r in top_value],
        'health': dict(health),
        'restock_urgency': [dict(r) for r in restock],
        'activity_trend': [dict(r) for r in trend],
    })

# ─────────────────────────────────────────────
# NEW: Export API
# ─────────────────────────────────────────────
@app.route('/api/export/csv')
def export_csv():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM items ORDER BY category, name').fetchall()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(['ID','Name','Category','Quantity','Unit','Low Stock Threshold','Price','Stock Value','Supplier','Notes','Last Updated','Status'])
    for r in rows:
        status = 'Out of Stock' if r['quantity']==0 else ('Low Stock' if r['quantity']<=r['low_stock'] else 'OK')
        writer.writerow([r['id'],r['name'],r['category'],r['quantity'],r['unit'],r['low_stock'],
                         r['price'],round(r['quantity']*r['price'],2),r['supplier'] or '',r['notes'] or '',r['updated_at'],status])
    buf.seek(0)
    return send_file(io.BytesIO(buf.getvalue().encode()),
                     mimetype='text/csv',
                     as_attachment=True,
                     download_name=f"inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.csv")

@app.route('/api/export/excel')
def export_excel():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM items ORDER BY category, name').fetchall()
        stats = conn.execute('SELECT COUNT(*) as t, SUM(quantity*price) as v, SUM(CASE WHEN quantity<=low_stock THEN 1 ELSE 0 END) as l FROM items').fetchone()

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Inventory"
    HDR_FILL  = PatternFill("solid", fgColor="1E2435")
    LOW_FILL  = PatternFill("solid", fgColor="3D2A00")
    OUT_FILL  = PatternFill("solid", fgColor="3D0000")
    OK_FILL   = PatternFill("solid", fgColor="0D2B20")
    HDR_FONT  = Font(name='Calibri', bold=True, color="A0B4D0", size=10)
    TITLE_FONT= Font(name='Calibri', bold=True, color="FFFFFF", size=14)
    thin = Side(style='thin', color="2A3045")
    border = Border(left=thin,right=thin,top=thin,bottom=thin)

    ws.merge_cells('A1:L1')
    ws['A1'] = f"StockWise Inventory Report — {datetime.now().strftime('%d %b %Y %H:%M')}"
    ws['A1'].font = TITLE_FONT
    ws['A1'].fill = PatternFill("solid", fgColor="0F1117")
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:L2')
    ws['A2'] = f"Total Items: {stats['t']}   |   Stock Value: ₹{round(stats['v'] or 0,2):,}   |   Low/Out of Stock: {stats['l']}"
    ws['A2'].font = Font(name='Calibri', color="6B7590", italic=True, size=9)
    ws['A2'].fill = PatternFill("solid", fgColor="181C27")
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 18

    headers = ['ID','Name','Category','Quantity','Unit','Low Stock ≤','Price (₹)','Stock Value','Supplier','Notes','Last Updated','Status']
    col_widths = [6,28,14,11,8,12,11,13,16,22,20,12]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[3].height = 22

    for ri, r in enumerate(rows, 4):
        status = 'Out of Stock' if r['quantity']==0 else ('Low Stock' if r['quantity']<=r['low_stock'] else 'OK')
        row_fill = OUT_FILL if status=='Out of Stock' else (LOW_FILL if status=='Low Stock' else OK_FILL)
        vals = [r['id'],r['name'],r['category'],r['quantity'],r['unit'],r['low_stock'],
                r['price'],round(r['quantity']*r['price'],2),r['supplier'] or '',r['notes'] or '',r['updated_at'],status]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.fill = row_fill; cell.border = border
            cell.font = Font(name='Calibri', color="E8ECF4", size=9)
            cell.alignment = Alignment(vertical='center',
                                       horizontal='center' if ci in [1,4,5,6,7,8,12] else 'left')
        ws.row_dimensions[ri].height = 18

    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f"A3:L{3+len(rows)}"

    ws2 = wb.create_sheet("By Category")
    ws2['A1'] = "Category Summary"
    ws2['A1'].font = Font(name='Calibri', bold=True, color="FFFFFF", size=13)
    ws2['A1'].fill = PatternFill("solid", fgColor="0F1117")
    ws2.merge_cells('A1:E1'); ws2.row_dimensions[1].height = 28

    for ci, h in enumerate(['Category','Items','Total Qty','Stock Value (₹)','Avg Price (₹)'],1):
        cell = ws2.cell(row=2, column=ci, value=h)
        cell.font = HDR_FONT; cell.fill = HDR_FILL; cell.border = border
        cell.alignment = Alignment(horizontal='center')
    ws2.column_dimensions['A'].width = 16; ws2.column_dimensions['B'].width = 8
    ws2.column_dimensions['C'].width = 12; ws2.column_dimensions['D'].width = 16; ws2.column_dimensions['E'].width = 14

    with get_db() as conn:
        cat_rows = conn.execute('SELECT category,COUNT(*) as cnt,SUM(quantity) as qty,ROUND(SUM(quantity*price),2) as val,ROUND(AVG(price),2) as avg_p FROM items GROUP BY category ORDER BY val DESC').fetchall()
    for ri, r in enumerate(cat_rows, 3):
        for ci, v in enumerate([r['category'],r['cnt'],r['qty'],r['val'],r['avg_p']],1):
            cell = ws2.cell(row=ri, column=ci, value=v)
            cell.font = Font(name='Calibri', color="E8ECF4", size=9)
            cell.fill = PatternFill("solid", fgColor="181C27"); cell.border = border
            cell.alignment = Alignment(horizontal='center' if ci>1 else 'left', vertical='center')
        ws2.row_dimensions[ri].height = 18

    ws3 = wb.create_sheet("Low Stock")
    ws3['A1'] = "Low Stock Alert Report"
    ws3['A1'].font = Font(name='Calibri', bold=True, color="F95F5F", size=13)
    ws3['A1'].fill = PatternFill("solid", fgColor="0F1117")
    ws3.merge_cells('A1:G1'); ws3.row_dimensions[1].height = 28

    for ci, h in enumerate(['Name','Category','Current Qty','Unit','Threshold','% of Threshold','Urgency'],1):
        cell = ws3.cell(row=2, column=ci, value=h)
        cell.font = HDR_FONT; cell.fill = HDR_FILL; cell.border = border
        cell.alignment = Alignment(horizontal='center')
    ws3.column_dimensions['A'].width = 28; ws3.column_dimensions['B'].width = 14
    ws3.column_dimensions['C'].width = 13; ws3.column_dimensions['D'].width = 8
    ws3.column_dimensions['E'].width = 12; ws3.column_dimensions['F'].width = 16; ws3.column_dimensions['G'].width = 12

    low_items = [r for r in rows if r['quantity'] <= r['low_stock']]
    low_items.sort(key=lambda x: x['quantity']/max(x['low_stock'],1))
    for ri, r in enumerate(low_items, 3):
        pct = round((r['quantity']/max(r['low_stock'],1))*100,1)
        urgency = 'CRITICAL' if r['quantity']==0 else ('High' if pct<50 else 'Medium')
        row_fill = OUT_FILL if r['quantity']==0 else LOW_FILL
        for ci, v in enumerate([r['name'],r['category'],r['quantity'],r['unit'],r['low_stock'],f"{pct}%",urgency],1):
            cell = ws3.cell(row=ri, column=ci, value=v)
            cell.font = Font(name='Calibri', color="E8ECF4", size=9, bold=(ci==7))
            cell.fill = row_fill; cell.border = border
            cell.alignment = Alignment(horizontal='center' if ci>1 else 'left', vertical='center')
        ws3.row_dimensions[ri].height = 18

    for sheet in [ws, ws2, ws3]:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "1E2435"

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f"inventory_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")

# ─────────────────────────────────────────────
# Browser notification config API
# ─────────────────────────────────────────────
@app.route('/api/alerts/config', methods=['GET'])
def get_alert_config():
    with get_db() as conn:
        row = conn.execute('SELECT enabled, last_notified FROM alert_config WHERE id=1').fetchone()
    return jsonify(dict(row) if row else {'enabled': 1, 'last_notified': None})

@app.route('/api/alerts/config', methods=['POST'])
def save_alert_config():
    data = request.get_json()
    with get_db() as conn:
        conn.execute('UPDATE alert_config SET enabled=? WHERE id=1',
                     (1 if data.get('enabled') else 0,))
        conn.commit()
    return jsonify({'message': 'Notification settings saved'})

@app.route('/api/alerts/low-stock')
def get_alert_low_stock():
    items = get_low_stock_items()
    for i in items:
        i['is_low'] = i['quantity'] <= i['low_stock']
        i['is_out'] = i['quantity'] == 0
    return jsonify({'count': len(items), 'items': items})

@app.route('/api/alerts/ack', methods=['POST'])
def ack_alert():
    now = datetime.now().isoformat(timespec='seconds')
    with get_db() as conn:
        conn.execute('UPDATE alert_config SET last_notified=? WHERE id=1', (now,))
        conn.commit()
    return jsonify({'last_notified': now})

# ─────────────────────────────────────────────
# SSE Boot Sequence
# ─────────────────────────────────────────────
def _sse(event, data):
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"

@app.route('/api/boot-stream')
def boot_stream():
    def generate():
        yield "retry: 0\n\n"
        yield _sse('step', {'msg': 'Checking Python runtime', 'status': 'running'})
        time.sleep(0.35)
        py_ver = f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
        yield _sse('step', {'msg': f'Python {py_ver} · {platform.system()} {platform.machine()}', 'status': 'ok'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Starting Flask application', 'status': 'running'})
        time.sleep(0.3)
        yield _sse('step', {'msg': f'Flask app mounted · debug={app.debug}', 'status': 'ok'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Locating SQLite database', 'status': 'running'})
        time.sleep(0.4)
        db_exists = os.path.exists(DB_PATH)
        db_size   = round(os.path.getsize(DB_PATH) / 1024, 1) if db_exists else 0
        yield _sse('step', {'msg': f'SQLite found · {os.path.basename(DB_PATH)} ({db_size} KB)', 'status': 'ok'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Verifying database schema', 'status': 'running'})
        time.sleep(0.35)
        try:
            with get_db() as conn:
                tables = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
            yield _sse('step', {'msg': f'Schema OK · tables: {", ".join(tables)}', 'status': 'ok'})
        except Exception as e:
            yield _sse('step', {'msg': f'Schema error: {e}', 'status': 'error'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Loading inventory records', 'status': 'running'})
        time.sleep(0.3)
        try:
            with get_db() as conn:
                item_count  = conn.execute('SELECT COUNT(*) FROM items').fetchone()[0]
                low_count   = conn.execute('SELECT COUNT(*) FROM items WHERE quantity<=low_stock').fetchone()[0]
                total_value = conn.execute('SELECT COALESCE(SUM(quantity*price),0) FROM items').fetchone()[0]
            yield _sse('step', {'msg': f'{item_count} items loaded · {low_count} low-stock · value ₹{round(total_value,2):,}', 'status': 'ok'})
        except Exception as e:
            yield _sse('step', {'msg': f'Inventory error: {e}', 'status': 'error'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Loading notification settings', 'status': 'running'})
        time.sleep(0.3)
        try:
            with get_db() as conn:
                cfg = conn.execute('SELECT enabled FROM alert_config WHERE id=1').fetchone()
                low_count = conn.execute('SELECT COUNT(*) FROM items WHERE quantity<=low_stock').fetchone()[0]
            if cfg and cfg['enabled']:
                yield _sse('step', {'msg': f'Browser notifications ON · {low_count} item(s) need attention', 'status': 'ok'})
            else:
                yield _sse('step', {'msg': 'Browser notifications disabled', 'status': 'ok'})
        except Exception as e:
            yield _sse('step', {'msg': f'Alert config error: {e}', 'status': 'error'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'Warming analytics engine', 'status': 'running'})
        time.sleep(0.4)
        try:
            with get_db() as conn:
                cat_count = conn.execute('SELECT COUNT(DISTINCT category) FROM items').fetchone()[0]
                log_count = conn.execute('SELECT COUNT(*) FROM activity_log').fetchone()[0]
            yield _sse('step', {'msg': f'Analytics ready · {cat_count} categories · {log_count} log entries', 'status': 'ok'})
        except Exception as e:
            yield _sse('step', {'msg': f'Analytics error: {e}', 'status': 'error'})
        time.sleep(0.2)
        yield _sse('step', {'msg': 'NeelKanth IMS · All systems operational', 'status': 'ok'})
        time.sleep(0.3)
        yield _sse('done', {'msg': 'Boot complete', 'timestamp': datetime.now().isoformat(timespec='seconds')})

    return Response(
        stream_with_context(generate()),
        mimetype='text/event-stream',
        headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no', 'Access-Control-Allow-Origin': '*'}
    )

# ─────────────────────────────────────────────
init_db()

if __name__ == '__main__':
    if sys.platform == 'win32':
        sys.stdout.reconfigure(encoding='utf-8')
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_DEBUG', '1') == '1'
    print("\n🏪  NeelKanth IMS — Inventory Management System")
    print("─" * 48)
    print(f"   Open: http://127.0.0.1:{port}")
    print("   Boot: SSE stream at /api/boot-stream")
    print("   Features: CRUD · Analytics · Export · Notifications")
    print("─" * 48)
    app.run(host='0.0.0.0', port=port, debug=debug, threaded=True)
